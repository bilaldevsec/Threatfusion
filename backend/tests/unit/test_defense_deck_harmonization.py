"""Unit tests for ThreatFusion Defense Presentation Deck Harmonizer (Milestone #26).

Verifies:
1. End-to-end execution of generate_harmonized_defense_deck for both synthetic fallback
   and template-backed generation.
2. Output presentation has exactly 15 slides.
3. Presence of all unrounded Rev2 metrics across slide text and speaker notes.
4. Confinement of legacy inverted-split numbers (89.77%, 94.95%, 82k, 175k) to slides 8 and 12
   with explicit retraction qualification.
5. Qualified Slide 14 wording: bounded recorded/local integration path verified, live Azure VM
   deployment open under TF-011, and strict absence of unqualified operational claims.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
import pytest

from scripts.generate_harmonized_defense_deck import (
    DEFAULT_INPUT_PATH,
    SPEAKER_NOTES_BY_SLIDE,
    generate_harmonized_defense_deck,
)

UNROUNDED_REV2_METRICS = (
    "0.1181361214680695",
    "212,009",
    "313",
    "95.266%",
    "92.818%",
    "94.026%",
    "8.8308x",
    "16.3958",
)

LEGACY_INVERTED_METRICS = (
    "89.77%",
    "94.95%",
    "82,332",
    "175,341",
)

SLIDE_14_QUALIFIED_BADGE = (
    "BOUNDED RECORDED/LOCAL INTEGRATION PATH VERIFIED (16/16 TEST GATES); "
    "LIVE AZURE/SIEM DEPLOYMENT OPEN (TF-011)"
)

SLIDE_14_QUALIFIED_FOOTER = (
    "Repository software spine (Suricata eve.json -> 14 features -> loopback mTLS 1.3 -> "
    "Frozen RF -> Correlation -> RFC1918 Threat-Intel -> Advisory -> Reversible SOAR) is "
    "verified by automated integration tests (test_slide14_suricata_live_ingestion_path.py); "
    "live Azure VM / production SIEM field deployment remains unverified and open under TF-011."
)


def _extract_slide_texts(prs: Any) -> list[str]:
    """Extract combined slide shape text, table text, and notes for each slide."""
    slide_texts: list[str] = []
    for slide in prs.slides:
        texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    texts.append(paragraph.text)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for paragraph in cell.text_frame.paragraphs:
                            texts.append(paragraph.text)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            for paragraph in slide.notes_slide.notes_text_frame.paragraphs:
                notes += paragraph.text + "\n"
        slide_texts.append("\n".join(texts) + "\nNOTES:\n" + notes)
    return slide_texts


def _verify_deck_invariants(prs: Any) -> None:
    """Verify common invariants on any harmonized presentation deck."""
    assert len(prs.slides) == 15, f"Expected 15 slides, found {len(prs.slides)}"

    slide_texts = _extract_slide_texts(prs)
    combined_deck_text = "\n".join(slide_texts)

    # 1. Assert all unrounded numbers appear in slide text or speaker notes
    for metric in UNROUNDED_REV2_METRICS:
        assert metric in combined_deck_text, f"Missing unrounded Rev2 metric: {metric}"

    # 2. Assert legacy numbers appear ONLY on slides 8 and 12, with retraction
    for s_idx, text in enumerate(slide_texts):
        slide_num = s_idx + 1
        for legacy in LEGACY_INVERTED_METRICS:
            if legacy in text:
                assert slide_num in (
                    8,
                    12,
                ), f"Legacy metric '{legacy}' appeared on forbidden slide {slide_num}"
                lower_text = text.lower()
                assert (
                    "retracted" in lower_text
                    or "split inversion" in lower_text
                    or "swapped" in lower_text
                    or "inverted" in lower_text
                ), f"Legacy metric '{legacy}' on slide {slide_num} lacks explicit retraction qualification"

    # 3. Assert Slide 14 qualified badge and footer
    slide_14_text = slide_texts[13]  # 0-indexed slide 13 is Slide 14
    assert (
        SLIDE_14_QUALIFIED_BADGE in slide_14_text
    ), "Slide 14 is missing the required qualified integration badge"
    assert (
        SLIDE_14_QUALIFIED_FOOTER in slide_14_text
    ), "Slide 14 is missing the required qualified software spine footer"

    # 4. Assert Slide 14 contains no unqualified operational claims
    assert (
        "VERIFIED & OPERATIONAL" not in slide_14_text
    ), "Slide 14 contains unqualified claim 'VERIFIED & OPERATIONAL'"
    assert (
        "100% operational" not in slide_14_text
    ), "Slide 14 contains unqualified claim '100% operational'"


def test_generate_harmonized_defense_deck_synthetic(tmp_path: Path) -> None:
    """Verify synthetic deck generation without requiring external template file."""
    output_path = tmp_path / "ThreatFusion_Synthetic_Defense.pptx"
    generated_path = generate_harmonized_defense_deck(
        input_path=None,
        output_path=output_path,
    )
    assert generated_path == output_path
    assert output_path.exists()
    assert output_path.stat().st_size > 0

    prs = Presentation(str(output_path))
    _verify_deck_invariants(prs)


def test_generate_harmonized_defense_deck_missing_template_fallback(tmp_path: Path) -> None:
    """Verify that a nonexistent template path falls back gracefully to synthetic deck."""
    nonexistent = tmp_path / "does_not_exist.pptx"
    output_path = tmp_path / "ThreatFusion_Fallback_Defense.pptx"
    generated_path = generate_harmonized_defense_deck(
        input_path=nonexistent,
        output_path=output_path,
    )
    assert generated_path == output_path
    assert output_path.exists()
    prs = Presentation(str(output_path))
    assert len(prs.slides) == 15


def test_generate_harmonized_defense_deck_template_backed(tmp_path: Path) -> None:
    """Verify template-backed deck generation when the draft PPTX template exists."""
    if not DEFAULT_INPUT_PATH.exists():
        pytest.skip(f"Template '{DEFAULT_INPUT_PATH}' not present on disk in this environment")

    output_path = tmp_path / "ThreatFusion_Template_Defense.pptx"
    generated_path = generate_harmonized_defense_deck(
        input_path=DEFAULT_INPUT_PATH,
        output_path=output_path,
    )
    assert generated_path == output_path
    assert output_path.exists()
    assert output_path.stat().st_size > 0

    prs = Presentation(str(output_path))
    _verify_deck_invariants(prs)


def test_speaker_notes_dictionary_integrity() -> None:
    """Verify that speaker notes are defined for all 15 slides with exact Rev2 content."""
    assert len(SPEAKER_NOTES_BY_SLIDE) == 15
    for slide_idx in range(1, 16):
        note = SPEAKER_NOTES_BY_SLIDE.get(slide_idx)
        assert note is not None and isinstance(note, str)
        assert len(note.strip()) > 50, f"Slide {slide_idx} speaker note is too short or empty"

    # Verify specific unrounded numbers in notes
    assert "0.1181361214680695" in SPEAKER_NOTES_BY_SLIDE[4]
    assert "212,009" in SPEAKER_NOTES_BY_SLIDE[2]
    assert "313" in SPEAKER_NOTES_BY_SLIDE[2]
    assert "8.8308x" in SPEAKER_NOTES_BY_SLIDE[11]
    assert "16.3958" in SPEAKER_NOTES_BY_SLIDE[11]
    assert "TF-011" in SPEAKER_NOTES_BY_SLIDE[14]
